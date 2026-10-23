"""Lecture du classeur en lecture seule, avec empreinte et références de cellules."""
from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from .schemas import SourceCell, SourceRef


def sha256_file(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def slug(text: str) -> str:
    t = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")


def default_field_id(header: str | None, col: int) -> str:
    """Identifiant fondé sur la position (stable même si deux colonnes ont le même intitulé)."""
    letter = get_column_letter(col)
    return f"col_{letter}" if not header or not str(header).strip() else f"{slug(str(header))}__{letter}"


@dataclass
class Workbook:
    path: Path
    sha256: str
    sheet: str
    headers: list[str | None]
    rows: dict[int, list[SourceCell]] = field(default_factory=dict)  # n° de ligne Excel -> cellules
    max_row: int = 0
    max_col: int = 0
    empty_rows: list[int] = field(default_factory=list)

    def cell(self, row: int, col: int) -> SourceCell:
        return self.rows[row][col - 1]

    @property
    def data_rows(self) -> list[int]:
        return sorted(self.rows)


def read_workbook(path: str | Path, sheet: str = "Feuil1", field_ids: dict[int, str] | None = None) -> Workbook:
    """Lit le classeur sans jamais le modifier. field_ids : colonne -> identifiant du dictionnaire."""
    path = Path(path)
    digest = sha256_file(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    headers = [str(h).strip() if h is not None and str(h).strip() else None for h in all_rows[0]]
    max_col = len(headers)
    out = Workbook(path=path, sha256=digest, sheet=sheet, headers=headers, max_row=len(all_rows), max_col=max_col)
    for r_idx, values in enumerate(all_rows[1:], start=2):
        values = list(values) + [None] * (max_col - len(values))
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            out.empty_rows.append(r_idx)
            continue
        cells = []
        for c_idx, v in enumerate(values[:max_col], start=1):
            fid = (field_ids or {}).get(c_idx) or default_field_id(headers[c_idx - 1], c_idx)
            ref = SourceRef(
                file_sha256=digest, sheet=sheet, row=r_idx, col=c_idx, col_letter=get_column_letter(c_idx),
                original_header=headers[c_idx - 1], field_id=fid,
            )
            cells.append(SourceCell(ref=ref, raw_value=v, raw_type=type(v).__name__))
        out.rows[r_idx] = cells
    # vérification d'intégrité : le fichier n'a pas changé pendant la lecture
    if sha256_file(path) != digest:
        raise RuntimeError("Le fichier source a changé pendant la lecture")
    return out
