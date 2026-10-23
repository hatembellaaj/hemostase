"""
Génère dictionary/fields.yaml (dictionnaire initial) à partir des en-têtes du classeur.
À exécuter une fois ; le YAML est ensuite édité à la main par les responsables du registre.

    python dictionary/build_initial_dictionary.py data/raw/registre.xlsx

Statuts :
  confirmed   définition validée par un responsable (champ `approved_by` obligatoire)
  hypothesis  proposition issue du profilage : JAMAIS utilisée comme règle clinique active
  missing     définition inconnue, à compléter
"""
import sys
from pathlib import Path

import openpyxl
import yaml
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent

UNKNOWN = ["?", "??", "???", "????", "np", "NP", "nf", "NF", "NF?"]

# Codes : correspondances proposées d'après le profilage (croisements de colonnes). Toutes à l'état d'hypothèse.
OUI_NON_NP = {1: "oui", 2: "non", 3: "non précisé"}

SCOPE = [
    # col, field_id, nom contextualisé, type, unité, extras
    (1, "patient_name", "Nom et initiale du patient", "identifier", None, {"role": "direct_identifier"}),
    (3, "city", "Ville de résidence", "identifier", None, {"role": "quasi_identifier"}),
    (4, "file_number", "Numéro de dossier hospitalier", "identifier", None, {"role": "direct_identifier"}),
    (5, "birth_date", "Date de naissance", "date", None, {"role": "quasi_identifier"}),
    (6, "age_recorded", "Âge enregistré (date de référence inconnue)", "integer", "an",
     {"definition_note": "Concorde avec ~2024 − année de naissance ; ne correspond PAS à l'âge exact à la date DDN (écart +1 le plus souvent). Date de référence à confirmer."}),
    (8, "sex", "Sexe", "code", None, {"codes": {1: "masculin", 2: "féminin"}}),
    (10, "consanguinity", "Consanguinité des parents", "code", None, {"codes": OUI_NON_NP}),
    (11, "consanguinity_degree", "Degré de parenté des parents consanguins", "text", None,
     {"applicability": {"field": "consanguinity", "code": 1}}),
    (12, "family_similar_cases", "Cas similaire dans la famille", "code", None, {"codes": OUI_NON_NP}),
    (15, "family_case_relation", "Lien de parenté du ou des cas similaires", "text", None,
     {"applicability": {"field": "family_similar_cases", "code": 1}}),
    (23, "age_first_signs", "Âge aux premiers signes", "duration", None, {}),
    (25, "age_at_diagnosis_text", "Âge au diagnostic (texte)", "duration", None, {}),
    (26, "age_at_diagnosis_months", "Âge au diagnostic (en mois)", "integer", "mois",
     {"definition_note": "Intitulé « mois » : concorde avec l'âge au diagnostic texte pour 79 lignes sur 82 (hypothèse)."}),
    (29, "aptt_patient_control", "TCA patient/témoin", "ratio", "s", {}),
    (30, "aptt_ratio", "Ratio TCA", "number", "sans unité",
     {"definition_note": "Hypothèse : patient / témoin du TCA, enregistré tronqué ou arrondi au centième."}),
    (31, "aptt_mix_patient_control", "TCA mélange (M+T) patient/témoin", "ratio", "s", {}),
    (35, "fviii_level", "Taux de facteur VIII", "number", "%",
     {"definition_note": "Unité % supposée (valeurs « 1.5% », « <1% »). Valeurs 0–1 sans unité : % ou UI/ml ? à confirmer."}),
    (36, "fix_level", "Taux de facteur IX", "number", "%", {}),
    (37, "clinical_form", "Forme clinique (hémophilie A/B)", "code", None, {"codes": {1: "A", 2: "B"}}),
    (38, "severity", "Sévérité", "code", None, {"codes": {1: "sévère", 2: "modérée", 3: "mineure"}}),
    (41, "hemarthrosis", "Hémarthrose (antécédent)", "code", None, {"codes": {1: "oui", 2: "non"}}),
    (42, "hemarthrosis_elbow_right", "Hémarthrose coude droit", "code", None, {"codes": {1: "oui", 2: "non"}}),
    (43, "hemarthrosis_elbow_left", "Hémarthrose coude gauche", "code", None, {"codes": {1: "oui", 2: "non"}}),
    (44, "hemarthrosis_elbow_any", "Hémarthrose coude (synthèse)", "code", None, {"codes": {1: "oui", 2: "non"}}),
    (79, "circumcision", "Circoncision", "code", None,
     {"codes": OUI_NON_NP, "text_values": {"pas encore": "non encore réalisée (variante lexicale proposée)"}}),
    (82, "circumcision_age", "Âge à la circoncision", "duration", None,
     {"applicability": {"field": "circumcision", "code": 1}}),
    (125, "inhibitor", "Inhibiteur", "code", None, {"codes": OUI_NON_NP}),
    (126, "inhibitor_responder", "Faible/fort répondeur", "code", None, {"codes": {1: "faible répondeur", 2: "fort répondeur"},
                                                                        "applicability": {"field": "inhibitor", "code": 1}}),
    (127, "inhibitor_titer", "Titre de l'inhibiteur", "number", None,
     {"definition_note": "Unité non indiquée (UB/ml probable) : à confirmer.", "applicability": {"field": "inhibitor", "code": 1}}),
    (162, "ddn", "DDN (sens à confirmer)", "date", None,
     {"role": "quasi_identifier", "definition_note": "Distincte de la date de naissance ; 82/83 valeurs en 2022–2024 : date des dernières nouvelles ? À confirmer."}),
    (163, "comment_1", "Commentaire libre 1 (colonne sans titre)", "narrative", None, {}),
    (164, "comment_2", "Commentaire libre 2 (colonne sans titre)", "narrative", None, {}),
]


def main(xlsx: str) -> None:
    ws = openpyxl.load_workbook(xlsx, read_only=True)["Feuil1"]
    headers = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    fields = []
    for col, fid, name, ftype, unit, extra in SCOPE:
        h = headers[col - 1]
        entry = {
            "field_id": fid,
            "source_col": col,
            "source_letter": get_column_letter(col),
            "original_header": h,
            "contextual_name": name,
            "type": ftype,
            "unit": unit,
            "unit_status": "hypothesis" if unit else "missing",
            "definition_status": "hypothesis" if ftype not in ("identifier",) else "confirmed",
            "approved_by": None,
            "approved_on": None,
            "unknown_markers": UNKNOWN if ftype not in ("identifier", "narrative") else [],
            "in_scope": ftype not in ("identifier",),
        }
        if extra.get("role"):
            entry["privacy_role"] = extra["role"]
        if extra.get("definition_note"):
            entry["definition_note"] = extra["definition_note"]
        if "codes" in extra:
            entry["allowed_values"] = sorted(extra["codes"])
            entry["code_map"] = {k: {"label": v, "status": "hypothesis"} for k, v in extra["codes"].items()}
        if "text_values" in extra:
            entry["text_variants"] = {k: {"meaning": v, "status": "hypothesis"} for k, v in extra["text_values"].items()}
        if "applicability" in extra:
            entry["applicability"] = {**extra["applicability"], "status": "hypothesis"}
        if ftype == "identifier":
            entry["definition_status"] = "confirmed"
            entry["approved_by"] = "privacy-default"
        fields.append(entry)
    out = {
        "version": "0.1.0",
        "source_sheet": "Feuil1",
        "note": "Dictionnaire initial généré par profilage. Les hypothèses doivent être confirmées par les "
        "responsables du registre avant toute utilisation comme règle active. Les colonnes non listées "
        "conservent un identifiant fondé sur leur position (col_<lettre>).",
        "fields": fields,
    }
    (ROOT / "fields.yaml").write_text(yaml.safe_dump(out, allow_unicode=True, sort_keys=False, width=120), encoding="utf-8")
    print(f"{len(fields)} champs -> {ROOT / 'fields.yaml'}")


if __name__ == "__main__":
    main(sys.argv[1])
